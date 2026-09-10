# database/repository.py
from __future__ import annotations
from typing import Optional, Tuple, List
from datetime import datetime

from sqlalchemy import select, func, distinct, text
from sqlalchemy.exc import IntegrityError

from database.session import create_session_factory, db_session
from database.models import Base, IP, RootDomain, ScanRun, ScanTarget, Domain, Area, Domain_ip, Otx_httpscan, Shodan_http, Shodan_product, Dns_zone, Dns_zone_ns, Domain_dns_zone, Domain_cname
from models.website import Website
from utils.logger import get_logger
from config import Config

logger = get_logger("database.manager.orm")

class DatabaseManagerORM:
    """以 SQLAlchemy ORM 實作的資料庫管理器"""

    def __init__(self, db_path: str):
        self.db_path = db_path
        self.engine, self.SessionFactory = create_session_factory(db_path)
        self._init_database()

    def _init_database(self) -> None:
        Base.metadata.create_all(self.engine)
        self._migrate_dns_zone_schema()
        
        logger.info(f"已建立/確認資料表（ORM）：{self.db_path}")
        # 自動匯入 area 資料（若存在檔案且表為空）
        try:
            csv_path = getattr(Config, 'CITY_DOMAIN_CSV_PATH', 'domain_of_city.csv')
            if csv_path:
                with db_session(self.SessionFactory) as session:
                    has_any = session.execute(select(func.count(Area.id))).scalar_one()
                import os
                if has_any == 0 and os.path.isfile(csv_path):
                    inserted = self.import_area_from_csv(csv_path)
                    logger.info(f"area 初始化匯入：{inserted} 筆（來源: {csv_path}）")
        except Exception as e:
            logger.warning(f"初始化匯入 area 失敗：{e}")

    def _migrate_dns_zone_schema(self) -> None:
        with self.engine.begin() as conn:
            table_names = {
                row[0]
                for row in conn.execute(
                    text("SELECT name FROM sqlite_master WHERE type='table'")
                ).fetchall()
            }
            if "dns_zone" not in table_names:
                return

            zone_cols = {
                row[1] for row in conn.execute(text("PRAGMA table_info(dns_zone)")).fetchall()
            }
            expected_zone_cols = {
                "id",
                "zone_apex",
                "soa_mname",
                "soa_rname",
                "status",
                "error_message",
                "zone_scope_kind",
                "zone_hosting_kind",
                "is_delegated",
                "checked_at",
                "created_at",
                "updated_at",
            }

            if zone_cols != expected_zone_cols:
                conn.execute(text("PRAGMA foreign_keys=OFF"))
                conn.execute(text("ALTER TABLE dns_zone RENAME TO dns_zone_old"))
                conn.execute(
                    text(
                        """
                        CREATE TABLE dns_zone (
                            id INTEGER NOT NULL PRIMARY KEY AUTOINCREMENT,
                            zone_apex VARCHAR NOT NULL,
                            soa_mname VARCHAR,
                            soa_rname VARCHAR,
                            status VARCHAR NOT NULL DEFAULT 'ok',
                            error_message VARCHAR,
                            zone_scope_kind VARCHAR,
                            zone_hosting_kind VARCHAR,
                            is_delegated INTEGER,
                            checked_at DATETIME NOT NULL,
                            created_at DATETIME NOT NULL,
                            updated_at DATETIME NOT NULL,
                            CONSTRAINT uq_dns_zone_apex UNIQUE (zone_apex)
                        )
                        """
                    )
                )
                conn.execute(
                    text(
                        """
                        INSERT INTO dns_zone (
                            id, zone_apex, soa_mname, soa_rname, status, error_message,
                            zone_scope_kind, zone_hosting_kind, is_delegated,
                            checked_at, created_at, updated_at
                        )
                        SELECT
                            id,
                            zone_apex,
                            soa_mname,
                            soa_rname,
                            status,
                            error_message,
                            NULL,
                            NULL,
                            NULL,
                            checked_at,
                            created_at,
                            updated_at
                        FROM dns_zone_old
                        """
                    )
                )
                conn.execute(text("DROP TABLE dns_zone_old"))
                conn.execute(text("CREATE INDEX idx_dns_zone_apex ON dns_zone (zone_apex)"))
                conn.execute(text("PRAGMA foreign_keys=ON"))

            ns_cols = {
                row[1] for row in conn.execute(text("PRAGMA table_info(dns_zone_ns)")).fetchall()
            }
            if "provider_kind" not in ns_cols:
                conn.execute(text("ALTER TABLE dns_zone_ns ADD COLUMN provider_kind VARCHAR"))
            if "provider_name" not in ns_cols:
                conn.execute(text("ALTER TABLE dns_zone_ns ADD COLUMN provider_name VARCHAR"))
            if "is_external" not in ns_cols:
                conn.execute(text("ALTER TABLE dns_zone_ns ADD COLUMN is_external INTEGER"))

        with db_session(self.SessionFactory) as session:
            self._refresh_all_dns_zone_classifications(session)

    def _extract_root_domain(self, hostname: str) -> str:
        parts = [p for p in (hostname or "").strip().lower().split(".") if p]
        if len(parts) >= 3 and parts[-2] == "edu" and parts[-1] == "tw":
            return ".".join(parts[-3:])
        if len(parts) >= 2:
            return ".".join(parts[-2:])
        return hostname.strip().lower()

    def _classify_ns_host(self, ns_host: str) -> tuple[str, str, int]:
        host = (ns_host or "").strip().lower()
        if not host:
            return "unknown", "unknown", 0
        if host.endswith(".ncku.edu.tw"):
            return "school_managed", "ncku", 0

        provider_rules = [
            ("cloudflare", "external_dns_provider", "cloudflare"),
            ("awsdns", "external_dns_provider", "aws_route53"),
            ("azure-dns", "external_dns_provider", "azure_dns"),
            ("dnspod", "external_dns_provider", "dnspod"),
            ("digitalocean", "external_dns_provider", "digitalocean"),
            ("gandi", "external_dns_provider", "gandi"),
            ("hinet", "telco_or_infra", "hinet"),
            ("seed.net", "telco_or_infra", "seednet"),
            ("cht.com.tw", "telco_or_infra", "cht"),
        ]
        for needle, kind, name in provider_rules:
            if needle in host:
                return kind, name, 1
        return "unknown", "unknown", 1

    def _refresh_dns_zone_classification(self, session, zone: Dns_zone) -> None:
        ns_records = session.execute(
            select(Dns_zone_ns).where(Dns_zone_ns.zone_id == zone.id)
        ).scalars().all()

        provider_kinds: set[str] = set()
        has_external = False
        for ns in ns_records:
            provider_kind, provider_name, is_external = self._classify_ns_host(ns.ns_host)
            ns.provider_kind = provider_kind
            ns.provider_name = provider_name
            ns.is_external = is_external
            provider_kinds.add(provider_kind)
            has_external = has_external or bool(is_external)

        root_domain = self._extract_root_domain(zone.zone_apex)
        is_delegated = 1 if zone.zone_apex != root_domain else 0
        zone.is_delegated = is_delegated
        zone.zone_scope_kind = "delegated_subzone" if is_delegated else "root_zone"

        if not ns_records:
            zone.zone_hosting_kind = "unknown"
        elif provider_kinds == {"school_managed"}:
            zone.zone_hosting_kind = "school_managed"
        elif provider_kinds and provider_kinds.issubset({"external_dns_provider", "telco_or_infra", "unknown"}):
            zone.zone_hosting_kind = "external_managed" if has_external else "unknown"
        elif "school_managed" in provider_kinds and has_external:
            zone.zone_hosting_kind = "mixed"
        else:
            zone.zone_hosting_kind = "unknown"

    def _refresh_all_dns_zone_classifications(self, session) -> None:
        zones = session.execute(select(Dns_zone)).scalars().all()
        for zone in zones:
            self._refresh_dns_zone_classification(session, zone)

    # ---------- 內部工具 ----------
    def _get_or_create_ip(self, session, address: str) -> IP:
        """
        獲取或創建 IP 記錄（單一 address 欄位）
        """
        if not address:
            raise ValueError("必須提供 IP 地址")
        
        # 查找現有IP記錄
        ip = session.execute(select(IP).where(IP.address == address)).scalar_one_or_none()
        
        if ip:
            return ip
        
        # 創建新IP記錄
        ip = IP(address=address)
        session.add(ip)
        session.flush()  # 取得自動編號 id
        return ip

    def _is_ipv6(self, address: str) -> bool:
        return ":" in address

    def _get_or_create_root_domain(self, session, root_domain_name: Optional[str], source: Optional[str] = None) -> Optional[RootDomain]:
        if not root_domain_name:
            return None
        root_domain = session.execute(select(RootDomain).where(RootDomain.name == root_domain_name)).scalar_one_or_none()
        if root_domain:
            return root_domain
        root_domain = RootDomain(name=root_domain_name, source=source)
        session.add(root_domain)
        session.flush()
        return root_domain

    def get_or_create_root_domain(self, root_domain_name: str) -> RootDomain:
        """
        獲取或創建根域名記錄（公有方法，供 domain_hierarchy 使用）
        """
        with db_session(self.SessionFactory) as session:
            return self._get_or_create_root_domain(session, root_domain_name)
    
    def get_root_domain_by_name(self, root_domain_name: str) -> Optional[RootDomain]:
        """
        根據名稱獲取根域名記錄
        """
        with db_session(self.SessionFactory) as session:
            return session.execute(
                select(RootDomain).where(RootDomain.name == root_domain_name)
            ).scalar_one_or_none()

    # ---------- 對外 API ----------
    def save_website(self, website: Website, root_domain_name: Optional[str] = None, when_latest_otx_checked: Optional[datetime] = None) -> bool:
        """
        保存網站資料（使用新的域名階層結構）
        
        注意：此方法假設域名階層已經通過 domain_hierarchy 模組建立
        支援多 IP 關聯（一個 domain 可對應多個 IP）
        """
        with db_session(self.SessionFactory) as session:
            try:
                # 讀取 Website 物件上暫存的 IP 列表（新格式）
                ipv4_list = getattr(website, "ipv4_list", []) or []
                ipv6_list = getattr(website, "ipv6_list", []) or []
                
                # 向下相容：支援舊的單一 IP 格式
                if not ipv4_list:
                    ipv4_val = getattr(website, "ipv4", None)
                    if ipv4_val:
                        ipv4_list = [ipv4_val]
                if not ipv6_list:
                    ipv6_val = getattr(website, "ipv6", None)
                    if ipv6_val:
                        ipv6_list = [ipv6_val]
                
                # 查找域名（使用新的階層結構）
                domain = self._find_domain_by_fqdn(session, website.fqdn)
                if not domain:
                    logger.warning(f"找不到域名記錄：{website.fqdn}，可能域名階層尚未建立")
                    return False
                
                root_domain = self._get_or_create_root_domain(session, root_domain_name)
                
                # 更新域名記錄
                if root_domain:
                    domain.root_id = root_domain.id
                domain.when_crawled = website.when_crawled
                domain.when_latest_otx_checked = when_latest_otx_checked
                
                # 更新 DNS 紀錄旗標（A / AAAA / CNAME）
                has_a = getattr(website, "has_a", None)
                has_aaaa = getattr(website, "has_aaaa", None)
                has_cname = getattr(website, "has_cname", None)
                if has_a is not None:
                    domain.a = has_a
                if has_aaaa is not None:
                    domain.aaaa = has_aaaa
                if has_cname is not None:
                    domain.cname = has_cname

                cname_targets = getattr(website, "cname_targets", []) or []
                for target in cname_targets:
                    self._upsert_domain_cname(session, domain.id, target)
                
                # 處理 IP 關聯（多對多）- 統一寫入單一 address 欄位
                saved_ips = []
                combined_ips = list(dict.fromkeys([*ipv4_list, *ipv6_list]))
                for addr in combined_ips:
                    ip = self._get_or_create_ip(session, address=addr)
                    self._create_domain_ip_relation(session, domain.id, ip.id)
                    saved_ips.append(addr)
                
                saved_ip_display = ",".join(saved_ips) if saved_ips else "(no IP)"
                
                logger.info(
                    f"更新域名：{website.fqdn} -> [{len(saved_ips)} IPs: {saved_ip_display[:50]}{'...' if len(saved_ip_display) > 50 else ''}] "
                    f"root_domain={root_domain_name or '-'} "
                    f"otx={'T' if when_latest_otx_checked else 'F'}"
                )
                return True
                
            except Exception as e:
                logger.error(f"保存域名失敗 {website.fqdn}: {e}")
                return False
    
    def _create_domain_ip_relation(self, session, domain_id: int, ip_id: int) -> None:
        """
        創建 domain_ip 關聯（如果不存在）
        """
        existing_relation = session.execute(
            select(Domain_ip).where(
                Domain_ip.domain_id == domain_id,
                Domain_ip.ip_id == ip_id
            )
        ).scalar_one_or_none()
        
        if not existing_relation:
            domain_ip_relation = Domain_ip(domain_id=domain_id, ip_id=ip_id)
            session.add(domain_ip_relation)

    def _upsert_domain_cname(self, session, domain_id: int, target: str) -> None:
        target_norm = (target or "").strip().lower().rstrip(".")
        if not target_norm:
            return

        existing = session.execute(
            select(Domain_cname).where(
                Domain_cname.domain_id == domain_id,
                Domain_cname.target == target_norm,
            )
        ).scalar_one_or_none()

        if existing:
            existing.checked_at = datetime.now()
            return

        session.add(
            Domain_cname(
                domain_id=domain_id,
                target=target_norm,
                checked_at=datetime.now(),
            )
        )

    def _upsert_dns_zone(self, session, zone_info: dict) -> Dns_zone:
        zone_apex = (zone_info.get("zone_apex") or "").strip().lower()
        zone = session.execute(
            select(Dns_zone).where(Dns_zone.zone_apex == zone_apex)
        ).scalar_one_or_none()

        if not zone:
            zone = Dns_zone(zone_apex=zone_apex)
            session.add(zone)
            session.flush()

        zone.soa_mname = zone_info.get("soa_mname")
        zone.soa_rname = zone_info.get("soa_rname")
        zone.status = zone_info.get("status") or "ok"
        zone.error_message = zone_info.get("error_message")
        zone.checked_at = datetime.now()

        # 以最新查詢結果覆蓋 NS 清單
        session.execute(
            text("DELETE FROM dns_zone_ns WHERE zone_id = :zone_id"),
            {"zone_id": zone.id},
        )

        seen_ns: set[str] = set()
        for ns in zone_info.get("ns_records", []) or []:
            ns_host = (ns.get("ns_host") or "").strip().lower().rstrip(".")
            if not ns_host or ns_host in seen_ns:
                continue
            seen_ns.add(ns_host)
            session.add(
                Dns_zone_ns(
                    zone_id=zone.id,
                    ns_host=ns_host,
                    ns_ttl=ns.get("ns_ttl"),
                    provider_kind=None,
                    provider_name=None,
                    is_external=None,
                    checked_at=datetime.now(),
                )
            )

        self._refresh_dns_zone_classification(session, zone)

        return zone

    def _create_domain_dns_zone_relation(self, session, domain_id: int, zone_id: int, matched_by: Optional[str]) -> None:
        existing = session.execute(
            select(Domain_dns_zone).where(
                Domain_dns_zone.domain_id == domain_id,
                Domain_dns_zone.zone_id == zone_id,
            )
        ).scalar_one_or_none()

        if existing:
            existing.matched_by = matched_by
            existing.checked_at = datetime.now()
            return

        session.add(
            Domain_dns_zone(
                domain_id=domain_id,
                zone_id=zone_id,
                matched_by=matched_by,
                checked_at=datetime.now(),
            )
        )

    def save_dns_zone_batch(self, zone_info_by_fqdn: dict[str, dict]) -> tuple[int, int]:
        """
        批次儲存 DNS zone 資訊。

        Returns:
            (zones_upserted, domain_zone_links)
        """
        zones_upserted = 0
        links_saved = 0
        if not zone_info_by_fqdn:
            return zones_upserted, links_saved

        with db_session(self.SessionFactory) as session:
            zone_obj_cache: dict[str, Dns_zone] = {}
            for fqdn, info in zone_info_by_fqdn.items():
                zone_apex = (info.get("zone_apex") or "").strip().lower()
                if not zone_apex:
                    continue

                domain = self._find_domain_by_fqdn(session, fqdn)
                if not domain:
                    logger.debug(f"[DNS Zone] 找不到 domain：{fqdn}，略過關聯")
                    continue

                zone = zone_obj_cache.get(zone_apex)
                if not zone:
                    zone = self._upsert_dns_zone(session, info)
                    zone_obj_cache[zone_apex] = zone
                    zones_upserted += 1
                self._create_domain_dns_zone_relation(session, domain.id, zone.id, info.get("matched_by"))
                links_saved += 1

        logger.info(f"[DNS Zone] 儲存完成：zone upsert {zones_upserted} 筆，domain-zone 關聯 {links_saved} 筆")
        return zones_upserted, links_saved
    
    def _find_domain_by_fqdn(self, session, fqdn: str) -> Optional[Domain]:
        """
        通過 FQDN 查找域名記錄（階層式存法）
        例如：health.mdic.ncku.edu.tw
        -> 從 tw 開始，逐層往下找 edu -> ncku -> mdic -> health
        """
        parts = fqdn.split('.')
        if not parts:
            return None
        
        # 從最右邊（根域名）開始查找
        current_domain = None
        
        # 先查找根域名（最後一部分，如 "tw"）
        root_label = parts[-1]
        current_domain = session.execute(
            select(Domain).where(
                Domain.leftmost_label == root_label,
                Domain.parent_domain_id.is_(None)
            )
        ).scalar_one_or_none()
        
        if not current_domain:
            return None
        
        # 如果只有一個部分（如"tw"），直接返回
        if len(parts) == 1:
            return current_domain
        
        # 從右到左逐層查找：tw -> edu -> ncku -> mdic -> health
        for i in range(len(parts) - 2, -1, -1):
            label = parts[i]
            current_domain = session.execute(
                select(Domain).where(
                    Domain.leftmost_label == label,
                    Domain.parent_domain_id == current_domain.id
                )
            ).scalar_one_or_none()
            
            if not current_domain:
                return None
        
        return current_domain

    def save_websites_batch(self, websites: List[Website], root_domain_name: Optional[str] = None, when_latest_otx_checked: Optional[datetime] = None) -> int:
        saved = 0
        for w in websites:
            if self.save_website(w, root_domain_name, when_latest_otx_checked):
                saved += 1
        return saved

    def get_website_by_fqdn(self, fqdn_name: str) -> Optional[Website]:
        """
        通過 FQDN 查找網站資料（使用新的域名階層結構）
        返回的 Website 物件包含所有關聯的 IP（ipv4_list, ipv6_list）
        """
        with db_session(self.SessionFactory) as session:
            # 查找域名
            domain = self._find_domain_by_fqdn(session, fqdn_name)
            if not domain:
                return None
            
            # 獲取所有相關的 IP 資料（透過 domain_ip 關聯）
            domain_ip_relations = session.execute(
                select(Domain_ip).where(Domain_ip.domain_id == domain.id)
            ).scalars().all()
            
            # 收集所有 IP（由單一 address 欄位拆分為 v4 / v6）
            ipv4_list = []
            ipv6_list = []
            
            for relation in domain_ip_relations:
                ip = session.execute(
                    select(IP).where(IP.id == relation.ip_id)
                ).scalar_one_or_none()
                if ip:
                    if self._is_ipv6(ip.address):
                        ipv6_list.append(ip.address)
                    else:
                        ipv4_list.append(ip.address)
            
            website = Website(
                fqdn=fqdn_name,
                url=None,
                protocol=None,
                status_code=None,
                redirect_to=None,
                title=None,
                when_crawled=domain.when_crawled if isinstance(domain.when_crawled, datetime) else datetime.fromisoformat(str(domain.when_crawled)) if domain.when_crawled else None
            )
            
            # 設定 IP 列表屬性（新格式）
            setattr(website, "ipv4_list", ipv4_list)
            setattr(website, "ipv6_list", ipv6_list)
            
            # 向下相容：設定單一 IP 屬性（取第一個）
            website.ipv4 = ipv4_list[0] if ipv4_list else None
            website.ipv6 = ipv6_list[0] if ipv6_list else None
            
            return website

    def get_statistics(self) -> Tuple[int, int]:
        """
        獲取統計資料（使用新的域名階層結構）
        """
        with db_session(self.SessionFactory) as session:
            # 總域名數量
            total = session.execute(select(func.count(Domain.id))).scalar_one()
            
            # 唯一 IP 數量（通過 domain_ip 關聯計算）
            uniq_ip = session.execute(
                select(func.count(distinct(Domain_ip.ip_id)))
            ).scalar_one()
            
            return total, uniq_ip

    # ---------- 匯入 Area CSV ----------
    def import_area_from_csv(self, csv_path: str, encoding: str = "utf-8") -> int:
        """從 CSV 匯入 area 資料（支援 sn,domain,city 或 name,domain 格式）"""
        import csv
        inserted = 0
        with db_session(self.SessionFactory) as session:
            with open(csv_path, "r", encoding=encoding, newline="") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    name = (row.get("name") or row.get("city") or "").strip()
                    domain = (row.get("domain") or "").strip().lower()
                    if not name or not domain:
                        continue
                    exists = session.execute(select(Area).where(Area.domain == domain)).scalar_one_or_none()
                    if exists:
                        continue
                    area = Area(name=name, domain=domain)
                    session.add(area)
                    inserted += 1
        logger.info(f"Area 匯入完成，新增 {inserted} 筆")
        return inserted

    def find_area_domain(self, domain: str) -> Optional[str]:
        """查詢 domain 是否在 area 表中，返回匹配的 area.domain"""
        with db_session(self.SessionFactory) as session:
            result = session.execute(
                select(Area.domain).where(Area.domain == domain)
            ).scalar_one_or_none()
            return result

    # ---------- 匯入 xlsx Root Domain ----------
    def import_root_domains_from_xlsx(self, xlsx_folder: str = "institution") -> int:
        """
        從指定資料夾的所有 xlsx 檔案匯入 root_domain 資料
        只匯入 .edu.tw 結尾的域名
        
        遵循與 Shodan 相同的 area 處理邏輯：
        - 如果 root domain 是 area（如 ntpc.edu.tw），則提取子域名（如 ytes.ntpc.edu.tw）
        - 如果不是 area（如 ntu.edu.tw），則直接使用 root domain
        
        Args:
            xlsx_folder: xlsx 檔案所在資料夾路徑
            
        Returns:
            成功匯入的筆數
        """
        import os
        from urllib.parse import urlparse
        
        try:
            import openpyxl
        except ImportError:
            logger.error("需要安裝 openpyxl 套件：pip install openpyxl")
            return 0
        
        if not os.path.isdir(xlsx_folder):
            logger.warning(f"xlsx 資料夾不存在：{xlsx_folder}")
            return 0
        
        # 收集所有 xlsx 檔案
        xlsx_files = [f for f in os.listdir(xlsx_folder) if f.endswith('.xlsx')]
        if not xlsx_files:
            logger.warning(f"資料夾 {xlsx_folder} 中沒有 xlsx 檔案")
            return 0
        
        logger.info(f"[xlsx] 開始從 {xlsx_folder} 匯入 root_domain，共 {len(xlsx_files)} 個檔案")
        
        # 先載入所有 area domain 用於比對
        area_domains = self._get_all_area_domains()
        logger.info(f"[xlsx] 載入 {len(area_domains)} 個 area domain 用於比對")
        
        all_targets: set = set()  # 改名：存放最終要匯入的目標（可能是 root domain 或子域名）
        loaded_files = 0
        skipped_files = 0
        
        for xlsx_file in xlsx_files:
            file_path = os.path.join(xlsx_folder, xlsx_file)
            wb = None
            try:
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                ws = wb.active
                file_targets: set = set()
                
                # 找到「網址」欄位的位置（從前5行中搜尋表頭）
                url_col_idx = None
                header_row = None
                for row_idx in range(1, 6):
                    for col_idx, cell in enumerate(ws[row_idx], 1):
                        if cell.value and str(cell.value).strip() == "網址":
                            url_col_idx = col_idx
                            header_row = row_idx
                            break
                    if url_col_idx:
                        break
                
                if not url_col_idx:
                    logger.warning(f"[xlsx] {xlsx_file} 找不到「網址」欄位，跳過")
                    skipped_files += 1
                    continue
                
                # 從表頭下一行開始讀取資料
                for row in ws.iter_rows(min_row=header_row + 1, min_col=url_col_idx, max_col=url_col_idx):
                    cell_value = row[0].value
                    if not cell_value:
                        continue
                    
                    url_str = str(cell_value).strip()
                    if not url_str:
                        continue
                    
                    # 解析 URL，根據 area 邏輯決定要匯入的目標
                    target = self._extract_target_from_url(url_str, area_domains)
                    if target and target.endswith('.edu.tw'):
                        all_targets.add(target)
                        file_targets.add(target)
                
                loaded_files += 1
                logger.info(f"[xlsx] {xlsx_file} 提取到 {len(file_targets)} 個不重複的 .edu.tw 目標")
                
            except Exception as e:
                skipped_files += 1
                logger.warning(f"[xlsx] 讀取 {xlsx_file} 失敗：{e}")
                continue
            finally:
                if wb is not None:
                    wb.close()
        
        logger.info(
            f"[xlsx] xlsx 檔案讀取完成：成功 {loaded_files}/{len(xlsx_files)}，"
            f"跳過 {skipped_files}，提取到 {len(all_targets)} 個不重複的 .edu.tw 目標"
        )
        
        # 寫入資料庫
        inserted = 0
        with db_session(self.SessionFactory) as session:
            for target_name in all_targets:
                exists = session.execute(
                    select(RootDomain).where(RootDomain.name == target_name)
                ).scalar_one_or_none()
                if exists:
                    continue
                root_domain = RootDomain(name=target_name, source="xlsx")
                session.add(root_domain)
                inserted += 1
        
        logger.info(f"[xlsx] Root domain 匯入完成，新增 {inserted} 筆")
        return inserted
    
    def _get_all_area_domains(self) -> set:
        """取得所有 area 的 domain 集合"""
        with db_session(self.SessionFactory) as session:
            result = session.execute(select(Area.domain)).scalars().all()
            return set(result)
    
    def get_all_area_domains(self) -> set:
        """取得所有 area 的 domain 集合（公開方法）"""
        return self._get_all_area_domains()
    
    def _extract_target_from_url(self, url_str: str, area_domains: set) -> Optional[str]:
        """
        從 URL 提取目標域名，遵循 area 處理邏輯
        
        - 如果 root domain 是 area（如 ntpc.edu.tw），提取子域名（如 ytes.ntpc.edu.tw）
        - 如果不是 area（如 ntu.edu.tw），直接返回 root domain
        
        Args:
            url_str: URL 字串
            area_domains: area domain 集合
            
        Returns:
            要匯入的目標域名
        """
        from urllib.parse import urlparse
        
        url_str = url_str.strip()
        
        # 如果沒有協定，加上 http://
        if not url_str.startswith(('http://', 'https://')):
            url_str = 'http://' + url_str
        
        try:
            parsed = urlparse(url_str)
            hostname = parsed.netloc or parsed.path.split('/')[0]
            hostname = hostname.lower().strip()
            
            if not hostname:
                return None
            
            # 去掉 www. 前綴
            if hostname.startswith('www.'):
                hostname = hostname[4:]
            
            # 驗證是否為有效域名
            if '.' not in hostname:
                return None
            
            # 對於 .edu.tw 結尾的域名，提取 xxx.edu.tw 部分
            parts = hostname.split('.')
            if len(parts) >= 3 and parts[-2] == 'edu' and parts[-1] == 'tw':
                # 取最後 3 段作為 root domain (如 ncku.edu.tw)
                root_domain = '.'.join(parts[-3:])
                
                # 檢查是否為 area domain
                if root_domain in area_domains:
                    # 是 area，提取子域名（類似 Shodan 的 _extract_subdomain_from_hostname）
                    # hostname 例如: ytes.ntpc.edu.tw 或 www.ytes.ntpc.edu.tw（已去掉 www）
                    # root_domain 例如: ntpc.edu.tw
                    # 要提取: ytes.ntpc.edu.tw
                    
                    if len(parts) > 3:
                        # 有子域名，取最接近 root domain 的那個
                        # parts = ['ytes', 'ntpc', 'edu', 'tw'] → 'ytes.ntpc.edu.tw'
                        subdomain_part = parts[-4]  # 取 root domain 前面那個
                        return f"{subdomain_part}.{root_domain}"
                    else:
                        # 沒有子域名（如 ntpc.edu.tw），直接返回 root domain
                        return root_domain
                else:
                    # 不是 area，直接返回 root domain
                    return root_domain
            
            return hostname
            
        except Exception:
            return None
    
    def _extract_root_domain_from_url(self, url_str: str) -> Optional[str]:
        """
        從 URL 字串提取 root domain
        例如: https://www.ncku.edu.tw/about/ → ncku.edu.tw
        """
        from urllib.parse import urlparse
        
        url_str = url_str.strip()
        
        # 如果沒有協定，加上 http://
        if not url_str.startswith(('http://', 'https://')):
            url_str = 'http://' + url_str
        
        try:
            parsed = urlparse(url_str)
            hostname = parsed.netloc or parsed.path.split('/')[0]
            hostname = hostname.lower().strip()
            
            if not hostname:
                return None
            
            # 去掉 www. 前綴
            if hostname.startswith('www.'):
                hostname = hostname[4:]
            
            # 驗證是否為有效域名
            if '.' not in hostname:
                return None
            
            # 對於 .edu.tw 結尾的域名，提取 xxx.edu.tw 部分
            parts = hostname.split('.')
            if len(parts) >= 3 and parts[-2] == 'edu' and parts[-1] == 'tw':
                # 取最後 3 段作為 root domain (如 ncku.edu.tw)
                return '.'.join(parts[-3:])
            
            return hostname
            
        except Exception:
            return None
    
    def get_all_root_domains(self, source: Optional[str] = None) -> List[str]:
        """
        取得所有 root_domain 的名稱列表
        
        Args:
            source: 過濾來源（'shodan' / 'xlsx' / None 表示全部）
            
        Returns:
            root domain 名稱列表
        """
        with db_session(self.SessionFactory) as session:
            query = select(RootDomain.name)
            if source:
                query = query.where(RootDomain.source == source)
            result = session.execute(query).scalars().all()
            return list(result)

    # ---------- Scan resume state ----------
    def create_scan_run(self, cert_pattern: str, mode: str) -> int:
        """建立一次掃描任務，回傳 scan_run.id。"""
        with db_session(self.SessionFactory) as session:
            scan_run = ScanRun(
                cert_pattern=cert_pattern,
                mode=mode,
                status="running",
                created_at=datetime.now(),
                updated_at=datetime.now(),
            )
            session.add(scan_run)
            session.flush()
            logger.info(f"[續跑] 建立 scan_run id={scan_run.id}, mode={mode}, cert_pattern={cert_pattern}")
            return scan_run.id

    def get_latest_resumable_scan_run(self, cert_pattern: str) -> Optional[tuple[int, str, str]]:
        """取得指定 cert_pattern 最近一筆未完成的掃描任務。"""
        with db_session(self.SessionFactory) as session:
            scan_run = session.execute(
                select(ScanRun)
                .where(
                    ScanRun.cert_pattern == cert_pattern,
                    ScanRun.status.in_(["running", "interrupted", "failed"]),
                )
                .order_by(ScanRun.created_at.desc())
                .limit(1)
            ).scalar_one_or_none()
            if not scan_run:
                return None
            return scan_run.id, scan_run.mode, scan_run.status

    def prepare_scan_run_for_resume(self, run_id: int) -> int:
        """
        將上次中斷時仍為 running 的 target 改回 pending，讓本次可重跑該 target。
        回傳被重置的 target 數量。
        """
        reset_count = 0
        with db_session(self.SessionFactory) as session:
            scan_run = session.get(ScanRun, run_id)
            if scan_run:
                scan_run.status = "running"
                scan_run.updated_at = datetime.now()

            running_targets = session.execute(
                select(ScanTarget).where(
                    ScanTarget.run_id == run_id,
                    ScanTarget.status == "running",
                )
            ).scalars().all()
            for target in running_targets:
                target.status = "pending"
                target.error_message = None
                target.updated_at = datetime.now()
                reset_count += 1
        return reset_count

    def create_scan_targets(
        self,
        run_id: int,
        targets: List[tuple[str, str | None]],
    ) -> int:
        """
        建立本次掃描的 target 清單。

        Args:
            run_id: scan_run.id
            targets: [(target, source)]，source 例如 shodan / xlsx
        """
        inserted = 0
        seen: set[str] = set()
        with db_session(self.SessionFactory) as session:
            for idx, (target, source) in enumerate(targets, start=1):
                target_name = (target or "").strip().lower().rstrip(".")
                if not target_name or target_name in seen:
                    continue
                seen.add(target_name)

                root_domain = self._get_or_create_root_domain(session, target_name, source=source)
                exists = session.execute(
                    select(ScanTarget).where(
                        ScanTarget.run_id == run_id,
                        ScanTarget.target == target_name,
                    )
                ).scalar_one_or_none()
                if exists:
                    continue

                session.add(
                    ScanTarget(
                        run_id=run_id,
                        root_domain_id=root_domain.id if root_domain else None,
                        target=target_name,
                        target_index=idx,
                        source=source,
                        status="pending",
                        created_at=datetime.now(),
                        updated_at=datetime.now(),
                    )
                )
                inserted += 1

        logger.info(f"[續跑] scan_run id={run_id} 建立 {inserted} 個 scan_target")
        return inserted

    def get_scan_targets(
        self,
        run_id: int,
        include_done: bool = False,
    ) -> List[tuple[int, int, str, str, Optional[str]]]:
        """
        取得 scan_run 的 target 清單。

        Returns:
            [(scan_target_id, target_index, target, status, source)]
        """
        with db_session(self.SessionFactory) as session:
            query = select(ScanTarget).where(ScanTarget.run_id == run_id)
            if not include_done:
                query = query.where(ScanTarget.status != "done")
            targets = session.execute(
                query.order_by(ScanTarget.target_index.asc())
            ).scalars().all()
            return [
                (target.id, target.target_index, target.target, target.status, target.source)
                for target in targets
            ]

    def mark_scan_target_running(self, scan_target_id: int) -> None:
        with db_session(self.SessionFactory) as session:
            target = session.get(ScanTarget, scan_target_id)
            if target:
                target.status = "running"
                target.started_at = target.started_at or datetime.now()
                target.error_message = None
                target.updated_at = datetime.now()

    def mark_scan_target_done(self, scan_target_id: int) -> None:
        with db_session(self.SessionFactory) as session:
            target = session.get(ScanTarget, scan_target_id)
            if target:
                target.status = "done"
                target.finished_at = datetime.now()
                target.updated_at = datetime.now()

    def mark_scan_target_failed(self, scan_target_id: int, error_message: str) -> None:
        with db_session(self.SessionFactory) as session:
            target = session.get(ScanTarget, scan_target_id)
            if target:
                target.status = "failed"
                target.error_message = (error_message or "")[:1000]
                target.finished_at = datetime.now()
                target.updated_at = datetime.now()

    def mark_scan_run_completed(self, run_id: int) -> None:
        with db_session(self.SessionFactory) as session:
            scan_run = session.get(ScanRun, run_id)
            if scan_run:
                scan_run.status = "completed"
                scan_run.finished_at = datetime.now()
                scan_run.updated_at = datetime.now()

    def mark_scan_run_failed(self, run_id: int, status: str = "failed") -> None:
        with db_session(self.SessionFactory) as session:
            scan_run = session.get(ScanRun, run_id)
            if scan_run:
                scan_run.status = status
                scan_run.updated_at = datetime.now()

    def import_shodan_vt_targets(self, vt_targets: List[str]) -> int:
        """
        將 Shodan 查詢到的 VT 目標批量匯入 root_domain 資料表
        
        Args:
            vt_targets: Shodan 查詢到的 VT 目標列表（已經過 area 處理）
            
        Returns:
            成功匯入的筆數
        """
        if not vt_targets:
            return 0
        
        logger.info(f"[Shodan] 開始匯入 {len(vt_targets)} 個 VT 目標到 root_domain")
        
        inserted = 0
        with db_session(self.SessionFactory) as session:
            for target in vt_targets:
                if not target or not target.endswith('.edu.tw'):
                    continue
                
                # 檢查是否已存在
                exists = session.execute(
                    select(RootDomain).where(RootDomain.name == target)
                ).scalar_one_or_none()
                
                if exists:
                    # 如果存在但 source 是 None，更新為 shodan
                    if exists.source is None:
                        exists.source = "shodan"
                    continue
                
                # 新增記錄
                root_domain = RootDomain(name=target, source="shodan")
                session.add(root_domain)
                inserted += 1
        
        logger.info(f"[Shodan] VT 目標匯入完成，新增 {inserted} 筆")
        return inserted

    # ---------- OTX HTTP Scans ----------
    # IoT 過濾規則列表：(欄位匹配, 關鍵字, 原因, 大小寫敏感)
    # 欄位匹配: "title" = 只檢查 Title 欄位, "body" = 只檢查 Body 欄位, "*" = 檢查所有欄位
    IOT_FILTER_RULES = [
        ("title", "Embedded Web Server", "Embedded Web Server", True),
        ("*", "Synology", "Synology NAS", True),
        ("body", "type= password", "Login Form", False),
        ("body", "goform", "Login Form", False),
    ]
    
    # IoT 類型優先級（數字越小優先級越高）
    IOT_PRIORITY = {
        "Synology NAS": 1,
        "Embedded Web Server": 2,
        "Grafana": 3,
        "VMware vCenter Server": 3,
        "Login Form": 4,
    }
    
    def _get_iot_priority(self, iot_type: str | None) -> int:
        """取得 IoT 類型的優先級，數字越小優先級越高"""
        if not iot_type:
            return 999  # NULL 表示非 IoT，優先級最低
        return self.IOT_PRIORITY.get(iot_type, 999)
    
    def _check_iot_device(self, http_scans: List[dict]) -> tuple[bool, str]:
        """
        檢查 http_scans 中是否有 IoT 設備特徵
        
        根據 IOT_FILTER_RULES 列表進行過濾判斷
        
        Args:
            http_scans: OTX API 返回的 http_scans 資料列表
        
        Returns:
            (True, 原因) 如果發現 IoT 設備，否則 (False, "")
        """
        for scan in http_scans:
            name = scan.get("name", "").lower()
            value = scan.get("value", "")
            
            # value 可能是數字，轉成字串
            if not isinstance(value, str):
                value = str(value)
            
            for field_match, keyword, reason, case_sensitive in self.IOT_FILTER_RULES:
                # 檢查欄位是否匹配
                if field_match != "*" and field_match not in name:
                    continue
                
                # 檢查關鍵字是否存在
                if case_sensitive:
                    if keyword in value:
                        return True, reason
                else:
                    if keyword.lower() in value.lower():
                        return True, reason
        
        return False, ""
    
    def save_http_scans(self, fqdn: str, http_scans: List[dict]) -> int:
        """
        儲存 OTX HTTP Scans 資料到 otx_httpscan 資料表
        
        如果 http_scans 包含 IoT 設備特徵，則標記該 domain 的 iot_type
        
        Args:
            fqdn: 完整域名
            http_scans: OTX API 返回的 http_scans 資料列表
                       每筆包含 {"name": "...", "value": "..."}
        
        Returns:
            成功寫入的筆數
        """
        if not http_scans:
            return 0
        
        # 檢查是否為 IoT 設備
        is_iot, iot_reason = self._check_iot_device(http_scans)
        
        with db_session(self.SessionFactory) as session:
            try:
                # 查找域名
                domain = self._find_domain_by_fqdn(session, fqdn)
                if not domain:
                    logger.warning(f"找不到域名記錄：{fqdn}，無法儲存 http_scans")
                    return 0
                
                # 如果檢測到 IoT，標記 iot_type（依優先級覆蓋）
                if is_iot:
                    new_priority = self._get_iot_priority(iot_reason)
                    old_priority = self._get_iot_priority(domain.type)
                    
                    if new_priority < old_priority:
                        domain.type = iot_reason
                        logger.info(f"[OTX IoT] {fqdn} 標記為 '{iot_reason}'")
                    elif domain.type:
                        logger.debug(f"[OTX IoT] {fqdn} 已有更高優先級標記 '{domain.type}'，保留原標記")
                
                inserted = 0
                for scan in http_scans:
                    name = scan.get("name", "")
                    value = scan.get("value", "")
                    
                    # value 可能是數字，轉成字串
                    if not isinstance(value, str):
                        value = str(value)
                    
                    if not name:
                        continue
                    
                    # 建立新記錄
                    record = Otx_httpscan(
                        domain_id=domain.id,
                        name=name,
                        value=value
                    )
                    session.add(record)
                    inserted += 1
                
                logger.info(f"儲存 {fqdn} 的 {inserted} 筆 http_scans 資料")
                return inserted
                
            except Exception as e:
                logger.error(f"儲存 {fqdn} http_scans 失敗: {e}")
                return 0

    def save_http_scans_batch(self, http_scans_dict: dict[str, List[dict]]) -> int:
        """
        批次儲存多個 FQDN 的 OTX HTTP Scans 資料
        
        Args:
            http_scans_dict: {fqdn: [http_scans_data]} 的字典
        
        Returns:
            成功寫入的總筆數
        """
        total_inserted = 0
        for fqdn, scans in http_scans_dict.items():
            inserted = self.save_http_scans(fqdn, scans)
            total_inserted += inserted
        
        logger.info(f"批次儲存 http_scans 完成，共 {total_inserted} 筆")
        return total_inserted

    # ---------- Shodan HTTP ----------
    # Shodan IoT 過濾規則（檢查 http.html 內容）
    SHODAN_IOT_RULES = [
        ("Synology", "Synology NAS", True),
        ("Embedded Web Server", "Embedded Web Server", True),
        ("type= password", "Login Form", False),
        ("type=\"password\"", "Login Form", False),
        ("type='password'", "Login Form", False),
    ]
    
    def _check_shodan_iot(self, html: str) -> tuple[bool, str]:
        """
        檢查 Shodan http.html 中是否有 IoT 設備特徵
        
        Args:
            html: Shodan http.html 內容
        
        Returns:
            (True, 原因) 如果發現 IoT 設備，否則 (False, "")
        """
        if not html:
            return False, ""
        
        for keyword, reason, case_sensitive in self.SHODAN_IOT_RULES:
            if case_sensitive:
                if keyword in html:
                    return True, reason
            else:
                if keyword.lower() in html.lower():
                    return True, reason
        
        return False, ""
    
    def save_shodan_http(self, fqdn: str, html: str) -> bool:
        """
        儲存 Shodan http.html 資料到 shodan_http 資料表
        並檢查是否為 IoT 設備，若是則標記 iot_type
        
        Args:
            fqdn: 完整域名
            html: Shodan http.html 內容
        
        Returns:
            是否成功儲存
        """
        with db_session(self.SessionFactory) as session:
            try:
                domain = self._find_domain_by_fqdn(session, fqdn)
                if not domain:
                    logger.warning(f"找不到域名記錄：{fqdn}，無法儲存 shodan_http")
                    return False
                
                # 儲存 html 內容
                record = Shodan_http(
                    domain_id=domain.id,
                    html=html
                )
                session.add(record)
                
                # 檢查是否為 IoT 設備
                is_iot, iot_reason = self._check_shodan_iot(html)
                if is_iot:
                    new_priority = self._get_iot_priority(iot_reason)
                    old_priority = self._get_iot_priority(domain.type)
                    
                    if new_priority < old_priority:
                        domain.type = iot_reason
                        logger.info(f"[Shodan IoT] {fqdn} 標記為 '{iot_reason}'")
                    elif domain.type:
                        logger.debug(f"[Shodan IoT] {fqdn} 已有更高優先級標記 '{domain.type}'，保留原標記")
                
                logger.debug(f"儲存 {fqdn} 的 shodan_http 資料")
                return True
                
            except Exception as e:
                logger.error(f"儲存 {fqdn} shodan_http 失敗: {e}")
                return False
    
    def save_shodan_http_batch(self, shodan_data: dict[str, str]) -> int:
        """
        批次儲存多個 FQDN 的 Shodan http.html 資料
        
        Args:
            shodan_data: {fqdn: html} 的字典
        
        Returns:
            成功儲存的筆數
        """
        saved_count = 0
        for fqdn, html in shodan_data.items():
            if self.save_shodan_http(fqdn, html):
                saved_count += 1
        
        logger.info(f"批次儲存 shodan_http 完成，共 {saved_count} 筆")
        return saved_count
    
    def mark_iot_type(self, fqdn: str, iot_type: str) -> bool:
        """
        手動標記 domain 的 IoT 類型（依優先級覆蓋）
        
        Args:
            fqdn: 完整域名
            iot_type: IoT 類型
        
        Returns:
            是否成功標記
        """
        with db_session(self.SessionFactory) as session:
            try:
                domain = self._find_domain_by_fqdn(session, fqdn)
                if not domain:
                    logger.warning(f"找不到域名記錄：{fqdn}")
                    return False
                
                new_priority = self._get_iot_priority(iot_type)
                old_priority = self._get_iot_priority(domain.type)
                
                if new_priority < old_priority:
                    domain.type = iot_type
                    logger.info(f"[IoT] {fqdn} 標記為 '{iot_type}'")
                    return True
                else:
                    logger.debug(f"[IoT] {fqdn} 已有更高優先級標記 '{domain.type}'")
                    return False
                    
            except Exception as e:
                logger.error(f"標記 {fqdn} 失敗: {e}")
                return False

    # ---------- Shodan Product ----------
    # Shodan Product IoT 過濾規則（檢查 product 欄位）
    SHODAN_PRODUCT_IOT_RULES = [
        ("Grafana", "Grafana"),
        ("VMware vCenter Server", "VMware vCenter Server"),
    ]
    
    def _check_shodan_product_iot(self, product: str) -> tuple[bool, str]:
        """
        檢查 Shodan product 欄位是否為 IoT 設備
        
        Args:
            product: Shodan product 欄位內容
        
        Returns:
            (True, 原因) 如果是 IoT 設備，否則 (False, "")
        """
        if not product:
            return False, ""
        
        for keyword, reason in self.SHODAN_PRODUCT_IOT_RULES:
            if keyword in product:
                return True, reason
        
        return False, ""
    
    def save_shodan_product(self, fqdn: str, product: str) -> bool:
        """
        儲存 Shodan product 資料到 shodan_product 資料表
        並檢查是否為 IoT 設備，若是則標記 iot_type
        
        Args:
            fqdn: 完整域名
            product: Shodan product 欄位內容
        
        Returns:
            是否成功儲存
        """
        with db_session(self.SessionFactory) as session:
            try:
                domain = self._find_domain_by_fqdn(session, fqdn)
                if not domain:
                    logger.warning(f"找不到域名記錄：{fqdn}，無法儲存 shodan_product")
                    return False
                
                # 儲存 product 內容
                record = Shodan_product(
                    domain_id=domain.id,
                    product=product
                )
                session.add(record)
                
                # 檢查是否為 IoT 設備
                is_iot, iot_reason = self._check_shodan_product_iot(product)
                if is_iot:
                    new_priority = self._get_iot_priority(iot_reason)
                    old_priority = self._get_iot_priority(domain.type)
                    
                    if new_priority < old_priority:
                        domain.type = iot_reason
                        logger.info(f"[Shodan Product] {fqdn} 標記為 '{iot_reason}'")
                    elif domain.type:
                        logger.debug(f"[Shodan Product] {fqdn} 已有更高優先級標記 '{domain.type}'，保留原標記")
                
                logger.debug(f"儲存 {fqdn} 的 shodan_product 資料: {product}")
                return True
                
            except Exception as e:
                logger.error(f"儲存 {fqdn} shodan_product 失敗: {e}")
                return False
    
    def save_shodan_product_batch(self, product_data: dict[str, list[str]]) -> int:
        """
        批次儲存多個 FQDN 的 Shodan product 資料
        
        Args:
            product_data: {fqdn: [product1, product2, ...]} 的字典
        
        Returns:
            成功儲存的筆數
        """
        saved_count = 0
        for fqdn, products in product_data.items():
            for product in products:
                if product and self.save_shodan_product(fqdn, product):
                    saved_count += 1
        
        logger.info(f"批次儲存 shodan_product 完成，共 {saved_count} 筆")
        return saved_count
